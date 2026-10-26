"""
補齊部件拆解（給 Claude 或熟悉程式的人執行，老師不用動）

新增課本（改了 shared/lessons.js）之後執行一次：
  shared/parts.js        只「新增」還沒有拆法的字（新字＋拆出來的部件再往下拆），原有內容一個字都不改
  docs/materials/parts-table.xlsx  新字加在最後面；原本就有的字，在「出現課次」補上新課本的課次；老師填的綠色欄位不動
再執行一次不會有任何變化（已經有的就跳過）。

需要：Python 3、openpyxl（pip install openpyxl）、Node.js
需要先下載（不放進 repo）：cjkvi-ids 的 ids.txt  https://github.com/cjkvi/cjkvi-ids（衍生自 CHISE，GPLv2 或更新版）
執行（在 repo 最上層）：
  python3 docs/scripts/make_parts.py <ids.txt>

拆法規則（和 2026-10-20 第一次產生 parts.js 時相同；拿當時的 1,256 個字重算，1,229 個完全一致）：
  1. 每個字優先用台灣字形（[T]）的拆法；沒有標地區就用第一種。
  2. 拆出來的部件如果是罕用字（擴充區 A 以後），再往下拆成常見部件。
  3. 資料庫裡打不出來的部分（帶圈數字、片假名等代用符號）略去，所以少數字只拆出一部分，部件表會標黃色提醒檢查。
"""
import copy, json, os, re, subprocess, sys
import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.worksheet.cell_range import MultiCellRange

IDS = sys.argv[1]
PARTS_JS, TABLE = 'shared/parts.js', 'docs/materials/parts-table.xlsx'

# 1. 讀 lessons.js 與 parts.js（用 Node 執行後轉成 JSON）
JS = r'''
const fs=require("fs");const o={};
new Function(require("./docs/scripts/lessons-src.js")()+";Object.assign(this,{BOOKS,PARTS})").call(o);
new Function(fs.readFileSync("shared/parts.js","utf8")+";this.PA=PARTS_AUTO").call(o);
console.log(JSON.stringify({BOOKS:o.BOOKS,PARTS:o.PARTS,PA:o.PA}));'''
L = json.loads(subprocess.run(['node', '-e', JS], capture_output=True, text=True, check=True).stdout)
BOOKS, PARTS, PA = L['BOOKS'], L['PARTS'], dict(L['PA'])
EDITION_ORDER = ['翰林', '南一', '康軒']
order = sorted(BOOKS, key=lambda b: (EDITION_ORDER.index(BOOKS[b]['edition']) if BOOKS[b]['edition'] in EDITION_ORDER else 99,
                                     BOOKS[b]['grade'] * 2 + (BOOKS[b]['term'] == '下')))

# 2. 讀 cjkvi-ids：字 → 拆法（優先台灣字形）
D = {}
for line in open(IDS, encoding='utf8'):
    if line.startswith('#'): continue
    f = line.rstrip('\n').split('\t')
    if len(f) < 3: continue
    pick = None
    for o in f[2:]:
        m = re.match(r'^(.*?)\[(\w+)\]$', o)
        if m and 'T' in m.group(2): pick = m.group(1); break
    if pick is None:
        plain = [o for o in f[2:] if '[' not in o]
        pick = plain[0] if plain else re.sub(r'\[\w+\]$', '', f[2])
    D[f[1]] = pick
IDC = set(chr(c) for c in range(0x2FF0, 0x2FFC))
STRUCT = {'⿰': '左右', '⿱': '上下', '⿲': '左中右', '⿳': '上中下', '⿴': '全包圍', '⿵': '上方包圍', '⿶': '下方包圍',
          '⿷': '左方包圍', '⿸': '左上包圍', '⿹': '右上包圍', '⿺': '左下包圍', '⿻': '交疊'}
STROKES = set('一丨丿丶乀乚𠃌乛亅𠄌㇀㇇㇉乁𡿨𠃊𠃋𠂆')
rare = lambda c: 0x3400 <= ord(c) <= 0x4DBF or ord(c) >= 0x20000
junk = lambda x: re.match(r'[①-⓿゠-ヿ　-〿]', x)

def first(c):
    s = D.get(c)
    if not s or s == c: return None
    t = [x for x in re.sub(r'&[^;]+;', '', s) if x not in IDC]
    return t if len(t) >= 2 or any(rare(x) for x in t) else None

def expand(c, depth=0):
    t = first(c)
    if t is None: return None
    out = []
    for x in t:
        e = expand(x, depth + 1) if rare(x) and depth < 6 else None
        out += e if e else [x]
    return [x for x in out if not junk(x)]

# 3. 補 parts.js：課本裡還沒有拆法的字，以及拆出來的部件再往下拆
chars = [c for b in order for n in BOOKS[b]['lessons'] for c in BOOKS[b]['lessons'][n]['write'] + BOOKS[b]['lessons'][n]['read']]
new, queue = {}, list(dict.fromkeys(chars))
while queue:
    c = queue.pop(0)
    if c in PA or c in new or c in PARTS: continue
    e = expand(c)
    if not e or e == [c]: continue
    new[c] = ' '.join(e)
    queue += [x for x in e if x not in PA and x not in new]
if new:
    s = open(PARTS_JS, encoding='utf8').read(); i = s.rindex('}')
    s = s[:i] + ',' + ','.join(f'{json.dumps(k, ensure_ascii=False)}:{json.dumps(v, ensure_ascii=False)}' for k, v in new.items()) + s[i:]
    open(PARTS_JS, 'w', encoding='utf8').write(s)
PA.update(new)

# 4. 補部件表
def parts(c): return (PARTS.get(c) or PA.get(c) or '').split()
def deep(c, seen=None):
    out = []
    for p in parts(c):
        if p in out: continue
        out.append(p)
        if p != c: out += [q for q in deep(p) if q not in out]
    return out
def flags(c):
    f, raw, ps = [], D.get(c, ''), parts(c)
    if re.search(r'&[^;]+;|[①-⓿゠-ヿ]', raw): f.append('只拆出部分部件')
    if len(ps) == 1: f.append('只有一個部件')
    if any(rare(p) for p in ps): f.append('含罕用部件，可能顯示成方框')
    if any(p in STROKES for p in ps): f.append('拆出單一筆畫，可能太細')
    return '；'.join(f) or None
def where(b, n, read):   # 例如「康軒5上第3課」；第 0 課（課前單元）用單元名稱，例如「康軒1上真簡單」
    les = BOOKS[b]['lessons'][n]
    return f"{BOOKS[b]['edition']}{BOOKS[b]['grade']}{BOOKS[b]['term']}" + (les.get('title') or '課前單元' if n == '0' else f'第{n}課') + ('(認讀)' if read else '')

wb = openpyxl.load_workbook(TABLE); ws = wb['生字部件']
rows = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1)}
yellow, green = PatternFill('solid', fgColor='FBF1DC'), PatternFill('solid', fgColor='E4F3EC')
added = updated = 0
for b in order:
    for n, les in BOOKS[b]['lessons'].items():
        for read, s in ((False, les['write']), (True, les['read'])):
            for c in s:
                w = where(b, n, read)
                if c in rows:
                    cell = ws.cell(rows[c], 2)
                    if w not in (cell.value or '').split('、'): cell.value = (cell.value + '、' if cell.value else '') + w; updated += 1
                    continue
                r, has = ws.max_row + 1, bool(parts(c))
                vals = [c, w, ' '.join(parts(c)) if has else None, (STRUCT.get(D.get(c, '')[:1]) if has else '獨體字'), None, None,
                        ' '.join(deep(c)) if has else None, ('人工（lessons.js）' if c in PARTS else '自動（parts.js）') if has else '—',
                        flags(c) if has else None, None, None, None]
                for j, v in enumerate(vals, 1):
                    cell, src = ws.cell(r, j, v), ws.cell(r - 1, j)
                    cell.font, cell.alignment, cell.border = copy.copy(src.font), copy.copy(src.alignment), copy.copy(src.border)
                    if j in (5, 6, 10, 11, 12): cell.fill = green
                if vals[8]: ws.cell(r, 9).fill = yellow
                rows[c] = r; added += 1
if added or updated:
    ws.auto_filter.ref = f'A1:L{ws.max_row}'
    for dv in ws.data_validations.dataValidation: dv.sqref = MultiCellRange(f'K2:K{ws.max_row}')
    wb.save(TABLE)
print(f'parts.js 新增 {len(new)} 個字；部件表新增 {added} 列、{updated} 個舊字補上出現課次')

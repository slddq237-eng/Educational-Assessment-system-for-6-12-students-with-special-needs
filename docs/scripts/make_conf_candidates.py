"""
干擾字候選清單（給 Claude 或熟悉程式的人執行，老師不用動）

有些字的干擾字（同音字＋形近字）少於 2 個，出題時程式只好用課文裡的其他字補。
這支程式幫這些字找「候選干擾字」，寫進 docs/materials/material-review.xlsx 的「干擾字候選」頁，請老師挑。
只是準備資料：不會改 shared/lessons.js，老師挑好後再由 Claude 改進教材資料。

候選字只從平台課本的生字裡找（學生比較可能學過，字型裡也一定有這個字），
同年級或更低年級的字排前面。
  同音字：和這個字在語詞裡的讀音完全一樣（聲調也一樣）；不夠的話列聲調不同的（標「音近」）
  形近字：有相同部件的字（往下拆到底）；越少見的部件、部件位置一樣的排前面
          （扌、艹、口這類很多字都有的部件分數很低，所以會排在後面）
  兩邊都出現的字標 ★（形音皆近）
排序只是方便老師挑，怎麼選、要選幾個，等規格書 A3 討論決定。

需要：Python 3、openpyxl（pip install openpyxl）、Node.js
需要先下載（不放進 repo）：字嗨注音標楷的讀音表 phonetic/phonic_table_Z.txt
  https://github.com/ButTaiwan/bpmfvs（和 make_bpmf.py 用同一份）
執行（在 repo 最上層）：
  python3 docs/scripts/make_conf_candidates.py <phonic_table_Z.txt>
"""
import json, subprocess, sys
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

XLSX = 'docs/materials/material-review.xlsx'
SHEET = '干擾字候選'
N_SOUND, N_SHAPE = 6, 6      # 每個字最多列幾個候選

data = json.loads(subprocess.run(['node', '-e', '''
const fs = require("fs");
const src = [require("./docs/scripts/lessons-src.js")(), ...["shared/parts.js", "shared/parts-layers.js", "shared/parts-api.js", "shared/bpmf.js"].map(f => fs.readFileSync(f, "utf8"))].join("\\n");
/* 部件和平台（部件提示、部件重組）一樣，透過 shared/parts-api.js 的 Parts 從第二層開始讀（人工手動 → 電腦自動，沒有再讀第三層），2026-10-31 步驟 3b */
new Function("window", src + ";const PT = Object.fromEntries([...new Set([...ALL_LESSON_CHARS, ...Object.keys(PARTS_AUTO)])].map(c => [c, (window.Parts.read(c, {from: 2}) || {parts: []}).parts])); process.stdout.write(JSON.stringify({BOOKS,WORDS,WORDS_ED,CONF,CONF_ED,PT,BPMF_CH}))")({});
'''], capture_output=True, text=True, check=True).stdout)
BOOKS, CONF, CONF_ED = data['BOOKS'], data['CONF'], data['CONF_ED']

READ = {}
for line in open(sys.argv[1], encoding='utf-8'):
    f = line.rstrip('\n').split('\t')
    if len(f) >= 4: READ[f[0]] = f[3:]

han = lambda c: '㐀' <= c <= '鿿'
def parts(c):
    return data['PT'].get(c, [])
def base(r): return r.rstrip('ˊˇˋ˙').lstrip('˙')
def edrank(e): return {'翰林': 0, '南一': 1, '康軒': 2}.get(e, 9)
ORDER = sorted(BOOKS, key=lambda b: (edrank(BOOKS[b]['edition']), BOOKS[b]['grade'] * 2 + (BOOKS[b]['term'] == '下')))
name = lambda b: f"{BOOKS[b]['edition']} {BOOKS[b]['grade']}{BOOKS[b]['term']}"

# 候選字池：所有課本的生字，記最低年級
POOL = {}
for b in BOOKS:
    for L in BOOKS[b]['lessons'].values():
        for c in (L.get('write') or '') + (L.get('read') or ''):
            if han(c): POOL[c] = min(POOL.get(c, 99), BOOKS[b]['grade'])

def reading(c, b):
    rs = READ.get(c) or ['']
    i = (data['BPMF_CH'].get(b) or {}).get(c, 0)
    return rs[i] if i < len(rs) else rs[0]

def sound_cands(c, b, grade, skip):
    r = reading(c, b)
    if not r: return [], r
    out = []
    for x, g in POOL.items():
        if x in skip: continue
        rx = (READ.get(x) or [''])[0]          # 候選字用它的預設讀音（學生最可能唸的音）
        if rx == r: out.append((0, g > grade, g, x, ''))
        elif base(rx) == base(r): out.append((1, g > grade, g, x, '音近'))
    out.sort()
    return [(x, tag) for _, _, _, x, tag in out[:N_SOUND]], r

STROKES = set('一丨丶丿乚乙亅𠃊𠃌㇀㇒𠄌')         # 單一筆畫不算部件
def comps(c, depth=0):
    """這個字的所有部件（往下拆到底）"""
    out = set()
    for p in parts(c):
        if p != c:
            out.add(p)
            if depth < 4: out |= comps(p, depth + 1)
    return out - STROKES
COMPS = {x: comps(x) for x in POOL}
FREQ = {}                                  # 每個部件出現在幾個字裡：越常見（像扌、艹、口）越不能代表字形相近
for x, cs in COMPS.items():
    for p in cs: FREQ[p] = FREQ.get(p, 0) + 1

def shape_cands(c, grade, skip):
    pc = parts(c)
    mine = comps(c) | {c}                  # 也找「包含這個字」的字（止 → 正、步）
    out = []
    for x, g in POOL.items():
        if x in skip: continue
        same = mine & (COMPS[x] | {x})
        same.discard(x)
        if not same: continue
        # 只留最大的共同部件（有了「咸」就不用再列「口」）
        same = {p for p in same if not any(p in COMPS.get(q, comps(q)) for q in same if q != p)}
        score = sum(1 / FREQ.get(p, 1) for p in same)
        px = parts(x)
        pos = sum(1 for i, p in enumerate(pc) if i < len(px) and px[i] == p and FREQ.get(p, 1) < 30)
        out.append((-score - pos, g > grade, g, x, '、'.join(sorted(same, key=lambda p: FREQ.get(p, 1))[:2])))
    out.sort()
    return [(x, s) for _, _, _, x, s in out[:N_SHAPE]]

rows = []
for b in ORDER:
    seen = set()
    for n, L in BOOKS[b]['lessons'].items():
        for kind, chars in (('習寫', L.get('write') or ''), ('認讀', L.get('read') or '')):
            for c in chars:
                if not han(c) or c in seen: continue
                seen.add(c)
                cf = (CONF_ED.get(b) or {}).get(c) or CONF.get(c) or ['', '']
                have = [x for x in (cf[0] or '') + (cf[1] or '') if han(x)]
                if len(have) >= 2: continue
                g = BOOKS[b]['grade']
                skip = set(have) | {c}
                snd, r = sound_cands(c, b, g, skip)
                shp = shape_cands(c, g, skip)
                both = {x for x, _ in snd} & {x for x, _ in shp}
                fmt_s = '　'.join(('★' if x in both else '') + x + (f'（{t}）' if t else '') for x, t in snd)
                fmt_p = '　'.join(('★' if x in both else '') + x + f'（{s}）' for x, s in shp)
                word = (data['WORDS_ED'].get(b) or {}).get(c) or data['WORDS'].get(c) or ''
                rows.append([name(b), n, kind, c, word, r, cf[0] or '', cf[1] or '', fmt_s or '（找不到）', fmt_p or '（找不到）', None, None])

wb = openpyxl.load_workbook(XLSX)
if SHEET in wb.sheetnames: del wb[SHEET]
ws = wb.create_sheet(SHEET, 1)
head = ['課本', '課次', '習寫／認讀', '字', '語詞', '讀音', '現有同音字', '現有形近字', '候選同音字', '候選形近字（相同部件）', '老師選用', '備註']
width = [11, 6, 9, 5, 12, 9, 10, 10, 34, 46, 16, 18]
ws.append(head)
hfill, gfill = PatternFill('solid', fgColor='1F3D36'), PatternFill('solid', fgColor='E2F0D9')
for i, c in enumerate(ws[1]):
    c.font, c.fill = Font(bold=True, color='FFFFFF'), hfill
    ws.column_dimensions[c.column_letter].width = width[i]
for r in rows: ws.append(r)
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(vertical='top', wrap_text=True)
    row[10].fill = row[11].fill = gfill
ws.freeze_panes = 'E2'
ws.auto_filter.ref = f'A1:L{len(rows) + 1}'

info = wb['說明']
for r in info.iter_rows():
    if r[0].value == '干擾字候選': break
else:
    info.append([])
    info.append(['干擾字候選', f'「干擾字候選」：干擾字少於 2 個的 {len(rows)} 個字，列出平台課本裡的同音字與形近字（相同部件）候選，★＝字形、字音都相近。'
                 '請在「老師選用」寫要用的字（例如「永勇」），不適合的就不用寫。候選字由程式自動找，只是方便挑選，不一定都適合。'])
wb.save(XLSX)
print(f'{len(rows)} 個字寫進 {XLSX}「{SHEET}」；找不到同音候選 {sum(r[8] == "（找不到）" for r in rows)} 個、形近候選 {sum(r[9] == "（找不到）" for r in rows)} 個')

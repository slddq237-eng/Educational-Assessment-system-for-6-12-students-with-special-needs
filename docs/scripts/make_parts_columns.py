"""
部件三層資料（給 Claude 或熟悉程式的人執行，老師不用動）

設計見 docs/design/parts-layers.md 與 parts-layers-impl.md。這支程式做三件事：
  1. 重算 docs/materials/parts-table.xlsx「生字部件」頁的電腦欄位（結構、第一～三層、正式評量遮第幾塊、需要檢查）；
     老師填的綠色欄位（老師修正：結構、各層人工手動與狀態、遮第幾塊人工手動、備註）一個字都不動。
  2. 產生 shared/parts-layers.js（平台之後透過 window.Parts 讀這個檔；步驟 3a 之前平台還不讀）。
  3. 過渡用：把狀態是「已確認」的人工手動拆法寫進 shared/lessons.js 的 PARTS_A、PARTS_B，讓平台目前的讀法不變
     （步驟 3a 平台改讀 parts-layers.js 之後，這一步和 lessons.js 的 PARTS_A、PARTS_B 會一起拿掉）。

三層（2026-10-31 維護者決定）：
  第一層  按結構拆成 2～3 個「可讀出聲音」的部件（部件是一個字，或是部首變體，見「部首變體對照」頁）
  第二層  可教學角度拆到最細：拆到部首就停，不拆成單一筆畫或罕用字；同一部件出現幾次寫幾次（森＝木 木 木）
  第三層  開源部件（cjkvi-ids）拆到最細，只存最底層（拼起來剛好是字，不含中間層）
每層讀取時「人工手動 → 電腦」；人工手動欄有填、狀態是已確認或待確認（空白當待確認）才算；狀態「待討論」不寫進平台。
正式評量遮擋（練習隨機）：左右、左中右遮最後一塊（右邊），上下、上中下遮最後一塊（下面），各種包圍遮裡面（第 2 塊）；
  交疊、獨體字、第一層拆不出來的不遮。老師可以在「正式評量遮第幾塊（人工手動）」填 1～3 改掉。
足、酋不再往下拆（2026-10-29 維護者決定）：每一層都當成一整塊，電腦把足字旁拆成「口 止」的合回「足」。

什麼時候跑：make_parts.py 之後（新增課本、改了部件表之後）。

需要：Python 3、openpyxl、Node.js
需要先下載（不放進 repo）：
  1. cjkvi-ids 的 ids.txt  https://github.com/cjkvi/cjkvi-ids（衍生自 CHISE，GPLv2 或更新版）
  2. pinyin-data 的 pinyin.txt  https://github.com/mozillazg/pinyin-data（MIT；讀音來自 Unicode Unihan 等）
     只拿來檢查「第一層的部件讀不讀得出聲音」，讀音本身不寫進表格或網站。
執行（在 repo 最上層）：
  python3 docs/scripts/make_parts_columns.py <ids.txt> <pinyin.txt>
"""
import copy, json, re, subprocess, sys, unicodedata
import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation

IDS, PINYIN = sys.argv[1], sys.argv[2]
TABLE, OUT, LES = 'docs/materials/parts-table.xlsx', 'shared/parts-layers.js', 'shared/lessons.js'
SHEET, VSHEET = '生字部件', '部首變體對照'
H = ['字', '出現課次', '結構', '老師修正：結構', '第一層（電腦）', '第一層（人工手動）', '第一層狀態', '第二層（電腦）', '第二層（人工手動）',
     '第二層狀態', '第三層（開源最細）', '正式評量遮第幾塊（電腦）', '正式評量遮第幾塊（人工手動）', '需要檢查', '備註']
COL = {h: i + 1 for i, h in enumerate(H)}
TEACHER = ['老師修正：結構', '第一層（人工手動）', '第一層狀態', '第二層（人工手動）', '第二層狀態', '正式評量遮第幾塊（人工手動）', '備註']
WIDTH = dict(zip(H, [6, 26, 10, 14, 16, 18, 12, 22, 20, 12, 22, 14, 16, 36, 24]))
STATUS = ['已確認', '待確認', '待討論']
STRUCTS = ['左右', '上下', '左中右', '上中下', '全包圍', '上方包圍', '下方包圍', '左方包圍', '左上包圍', '右上包圍', '左下包圍', '交疊', '獨體字']

# 1. 讀 parts.js 的 PARTS_AUTO（用 Node 執行後轉成 JSON）
js = '''const o={};new Function(require("fs").readFileSync("shared/parts.js","utf8")+";this.A=PARTS_AUTO").call(o);console.log(JSON.stringify(o.A));'''
AUTO = json.loads(subprocess.run(['node', '-e', js], capture_output=True, text=True, check=True).stdout)

# 2. 讀 cjkvi-ids（和 make_parts.py 一樣，優先台灣字形）
D = {}
for line in open(IDS, encoding='utf8'):
    if line.startswith('#'): continue
    f = line.rstrip('\n').split('\t')
    if len(f) < 3: continue
    pick = None
    for o in f[2:]:
        m = re.match(r'^(.*?)\[(\w+)\]$', o)
        if m and 'T' in m.group(2): pick = m.group(1); break
    if pick is None:
        plain = [o for o in f[2:] if '[' not in o]
        pick = plain[0] if plain else re.sub(r'\[\w+\]$', '', f[2])
    D[f[1]] = re.sub(r'&[^;]+;', '？', pick)
REV = {v: k for k, v in D.items() if len(v) > 1}          # 拆法 → 字（一塊是一小段組合時，找有沒有對應的字）
IDC = set(chr(c) for c in range(0x2FF0, 0x2FFC))
STRUCT = {'⿰': '左右', '⿱': '上下', '⿲': '左中右', '⿳': '上中下', '⿴': '全包圍', '⿵': '上方包圍', '⿶': '下方包圍',
          '⿷': '左方包圍', '⿸': '左上包圍', '⿹': '右上包圍', '⿺': '左下包圍', '⿻': '交疊'}
STROKES = set('一丨丿丶乀乚𠃌乛亅𠄌㇀㇇㇉乁𡿨𠃊𠃋𠂆')
rare = lambda c: 0x3400 <= ord(c) <= 0x4DBF or ord(c) >= 0x20000
junk = lambda x: re.match(r'[①-⓿゠-ヿ　-〿？]', x)
norm = lambda p: unicodedata.normalize('NFKC', p)

# 3. 讀音（只檢查有沒有讀音，不看是什麼音）
HAS_READING = set()
for line in open(PINYIN, encoding='utf8'):
    m = re.match(r'U\+([0-9A-F]+):\s*([^#]+)', line)
    if m and m.group(2).strip(): HAS_READING.add(chr(int(m.group(1), 16)))

# 4. 部首變體對照（老師可以在 Excel「部首變體對照」頁修改；第一層的部件是部首變體時，用對照的字判斷讀不讀得出聲音）
VARIANT_SEED = [('氵', '水', ''), ('扌', '手', ''), ('艹', '草', ''), ('亻', '人', ''), ('忄', '心', ''), ('犭', '犬', ''), ('礻', '示', ''),
                ('衤', '衣', ''), ('糹', '糸', ''), ('釒', '金', ''), ('⺮', '竹', ''), ('灬', '火', ''), ('刂', '刀', ''),
                ('阝', '阜／邑', '左邊是阜、右邊是邑，看位置'), ('辶', '辵', ''), ('攵', '攴', ''), ('廴', '廴', '本身就有讀音'),
                ('爫', '爪', ''), ('罒', '网', ''), ('牜', '牛', ''), ('飠', '食', ''), ('⺼', '肉', ''), ('⺈', '刀', ''), ('𧘇', '衣', ''), ('⺊', '卜', '')]
wb = openpyxl.load_workbook(TABLE)
if SHEET not in wb.sheetnames or [c.value for c in wb[SHEET][1]] != H:
    sys.exit('部件表欄位不是三層版本（見 docs/design/parts-layers-impl.md 步驟 2），先轉換欄位')
ws = wb[SHEET]
if VSHEET not in wb.sheetnames:
    vs = wb.create_sheet(VSHEET)
    vs.append(['部件', '對照成（讀音來源字）', '說明', '老師確認'])
    for r in VARIANT_SEED: vs.append(list(r))
    for j, w in enumerate([8, 22, 30, 14], 1): vs.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
    for c in vs[1]: c.font = copy.copy(ws.cell(1, 1).font); c.fill = copy.copy(ws.cell(1, 1).fill)
vs = wb[VSHEET]
chars_of = lambda v: [c for c in str(v) if not re.match(r'[／/、,，;；\s]', c)]   # 一格裡的每個字元算一個，分隔符號和空白略過
VARIANT = {}
for r in range(2, vs.max_row + 1):
    p, t = vs.cell(r, 1).value, vs.cell(r, 2).value
    if p and t:   # 「部件」「對照成」兩欄同一格都可以放多個：可以連寫（忄⺗）或用頓號、斜線、逗號、空格分開（忄、⺗）；每個字元算一個部件
        for q in chars_of(p): VARIANT[norm(q)] = chars_of(t)
readable = lambda p: norm(p) in HAS_READING or any(x in HAS_READING for x in VARIANT.get(norm(p), []))

# 5. 規則
KEEP_WHOLE = {'足', '酋'}                       # 和 shared/common.js 的 KEEP_WHOLE 一致
def keep_whole(v): return re.sub(r'^口 止(?= |$)', '足', v)
def struct_auto(c):
    s = D.get(c)
    return STRUCT.get(s[0], '獨體字') if s and s != c and s[0] in IDC else '獨體字'
def end(s, i):   # 從第 i 個字開始的「一塊」到哪裡結束（一塊＝一個字，或一段組合符號＋它的各塊）
    if i >= len(s): raise ValueError
    if s[i] not in IDC: return i + 1
    j = i + 1
    for _ in range(3 if s[i] in '⿲⿳' else 2): j = end(s, j)
    return j
def pieces(s):   # IDS 最上層切成兩塊或三塊
    if not s or s[0] not in IDC: return None
    try:
        out, i = [], 1
        for _ in range(3 if s[0] in '⿲⿳' else 2):
            j = end(s, i); out.append(s[i:j]); i = j
        return out if i == len(s) else None
    except ValueError:
        return None
def l1_auto(c):   # 回傳 (部件清單 or None, 原因)
    if c in KEEP_WHOLE: return None, ''
    s = D.get(c)
    if not s or s == c or s[0] not in IDC: return None, ''
    ps = pieces(s)
    if ps is None: return None, '第一層：資料格式特殊，拆不出 2～3 塊'
    out = []
    for p in ps:
        ch = p if len(p) == 1 and p not in IDC else REV.get(p)
        if ch is None or junk(ch): return None, '第一層：有一塊打不出來'
        out.append('足' if ch in ('𧾷', '⻊') else ch)
    bad = [p for p in out if not readable(p)]
    if bad: return None, '第一層：讀不出聲音的部件 ' + '、'.join(bad)
    return out, ''
def parts_auto(c):
    if c in KEEP_WHOLE: return []
    return [p for p in keep_whole(AUTO.get(c) or '').split() if p != c]
ok = lambda p: p not in STROKES and not rare(p) and not junk(p)
RADICAL = {unicodedata.normalize('NFKC', chr(c)) for c in range(0x2F00, 0x2FD6)} | \
          {unicodedata.normalize('NFKC', chr(c)) for c in range(0x2E80, 0x2EF4)} | \
          set('氵扌亻忄艹⺮糹訁釒阝辶⻊⺼礻衤灬爫罒⺈龶⺌⺍⺊⺗⺝⻌⻏⻖丬犭飠钅')
def fine(p, depth=0):   # 第二層：拆到部首就停，一個部件拆開後每一塊都是「看得懂的部件」才拆
    sub = parts_auto(p)
    if p not in RADICAL and depth < 8 and len(sub) >= 2 and all(ok(s) for s in sub):
        return [q for s in sub for q in fine(s, depth + 1)]
    return [p]
def l2_auto(c):
    first = parts_auto(c)
    return [q for p in first for q in fine(p)] if first else None
def leaves(p, depth=0):   # 第三層：一路拆到最底層
    sub = parts_auto(p)
    return [q for s in sub for q in leaves(s, depth + 1)] if depth < 8 and sub else [p]
def l3_auto(c):
    first = parts_auto(c)
    return [q for p in first for q in leaves(p, 1)] if first else None
def mask_auto(struct, l1):   # 正式評量遮第幾塊（1 起算）；沒有規則的回傳 None
    n = len(l1) if l1 else 0
    if struct in ('左右', '上下') and n == 2: return 2
    if struct in ('左中右', '上中下') and n == 3: return 3
    if '包圍' in struct and n == 2: return 2
    return None

# 6. 寫部件表、收集資料
yellow, green, none, red = (PatternFill('solid', fgColor='FBF1DC'), PatternFill('solid', fgColor='E4F3EC'), PatternFill(),
                            PatternFill('solid', fgColor='F4C7C3'))
note = {'結構': '電腦從 cjkvi-ids 判斷第一層的排列方式；不對請寫在右邊「老師修正：結構」',
        '老師修正：結構': '老師填：從下拉選單選；有填就用這個。遮擋題會用結構決定遮哪一塊',
        '第一層（電腦）': '電腦自動：按結構拆成 2～3 個可讀出聲音的部件，黃色＝待老師確認；拆不出來的留空，需要檢查欄會寫原因',
        '第一層（人工手動）': '老師填：部件用半形空格分開，順序照結構（左右＝左在前；上下＝上在前；包圍＝外層在前、裡面在後），例如「氵 青」。程式重算時不會動',
        '第一層狀態': '老師填（下拉選單）：已確認／待確認／待討論。人工手動欄有填、狀態空白當待確認；待討論不寫進平台',
        '第二層（電腦）': '電腦自動：可教學角度拆到最細（拆到部首就停，不拆成單一筆畫或罕用字），黃色＝待老師確認',
        '第二層（人工手動）': '老師填：部件用半形空格分開，例如「足 耳 耳 耳」；想當成獨體字不拆就寫字自己。程式重算時不會動',
        '第二層狀態': '老師填（下拉選單）：已確認／待確認／待討論。康軒四上 86 字舊的人工部件已搬到這裡、標待確認',
        '第三層（開源最細）': '電腦自動：cjkvi-ids 一路拆到最底層，拼起來剛好是字，不含中間層（含單一筆畫）；不用修改',
        '正式評量遮第幾塊（電腦）': '電腦自動：左右、上下遮右邊、下面那塊；左中右、上中下遮最後一塊；各種包圍遮裡面（第 2 塊）；沒有規則的留空。按「第一層」的順序數',
        '正式評量遮第幾塊（人工手動）': '老師填 1～3：優先對照「第一層」（第 1 塊＝第一層的第一個部件，順序照結構）。筆順不是先左後右、先上後下的字，或想改遮另一塊時填。要先有第一層（電腦或人工手動）才能填，沒有第一層的字請先補第一層，不會退到第二層；程式重算時不會動',
        '需要檢查': '電腦自動：可能有問題的地方，建議優先看這些列'}
sv = {}   # 字 → 資料
bad, stat = [], {'第一層': 0, '第二層': 0, '第三層': 0, '遮': 0}
unread, unread_ex = {}, {}
last = ws.cell(1, 1)
for j, h in enumerate(H, 1):
    c = ws.cell(1, j); c.font, c.fill, c.alignment, c.border = copy.copy(last.font), copy.copy(last.fill), copy.copy(last.alignment), copy.copy(last.border)
    ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = WIDTH[h]
    if h in note: c.comment = Comment(note[h], 'Claude')
def tokens(v): return [t for t in re.split(r'[\s＋+、,，]+', str(v).strip()) if t] if v not in (None, '') else []
def status_of(cell, has): return (str(cell).strip() if cell not in (None, '') else '待確認') if has else (str(cell).strip() if cell else '')
for r in range(2, ws.max_row + 1):
    c = ws.cell(r, 1).value
    if not c: continue
    g = lambda h: ws.cell(r, COL[h]).value
    reasons = []
    s_auto = struct_auto(c); s_fix = g('老師修正：結構'); struct = s_fix if s_fix in STRUCTS else s_auto
    if s_fix not in (None, '') and s_fix not in STRUCTS: bad.append(f'{c}（老師修正：結構）：「{s_fix}」不在選單裡'); reasons.append('老師修正：結構格式不對，沒有用')
    l1, why = l1_auto(c)
    if why:
        reasons.append(why)
        for p in re.findall(r'讀不出聲音的部件 (.*)', why)[0].split('、') if '讀不出聲音' in why else []:
            unread[p] = unread.get(p, 0) + 1; unread_ex.setdefault(p, []).append(c)
    l2, l3 = l2_auto(c), l3_auto(c)
    if l2 and any(rare(p) for p in l2): reasons.append('第二層含罕用部件，可能顯示成方框')
    if l2 and any(p in STROKES for p in l2): reasons.append('第二層含單一筆畫，可能太細')
    data = {'s': struct}
    layers = {}
    for n, key, auto, strict in ((1, '第一層', l1, True), (2, '第二層', l2, False)):
        man, stv = g(f'{key}（人工手動）'), g(f'{key}狀態')
        has = man not in (None, '')
        st = status_of(stv, has)
        if st and st not in STATUS: bad.append(f'{c}（{key}狀態）：「{stv}」不是已確認／待確認／待討論'); st = ''
        t = tokens(man) if has else []
        valid = has and st != '待討論' and t and all(len(x) == 1 for x in t) and (
            (len(t) in (2, 3)) if strict else (len(t) >= 1 and (len(t) > 1 or t[0] == c)))
        if has and not valid:
            if st != '待討論': bad.append(f'{c}（{key}人工手動）：「{man}」格式不對（每塊一個字，用空格分開；第一層要 2～3 塊）'); reasons.append(f'{key}人工手動格式不對，沒有寫進平台')
        d = {}
        if auto: d['a'] = ' '.join(auto)
        if valid: d['m'] = ' '.join(t)
        if st and (valid or st == '待討論'): d['t'] = st
        if d: layers[n] = d
        layers[(n, 'eff')] = t if valid else auto
    if l3: layers[3] = {'a': ' '.join(l3)}
    for n in (1, 2, 3):
        if n in layers: data[str(n)] = layers[n]
    eff1 = layers[(1, 'eff')]
    mk_auto = mask_auto(struct, l1 or eff1) if (l1 or eff1) else None
    mk_man = g('正式評量遮第幾塊（人工手動）')
    mk = mk_auto
    if mk_man not in (None, ''):
        try:
            n = int(mk_man)
            assert 1 <= n <= (len(eff1) if eff1 else 0)
            mk = n
        except (ValueError, AssertionError):
            bad.append(f'{c}（正式評量遮第幾塊（人工手動））：「{mk_man}」要填 1～{len(eff1) if eff1 else 0} 的整數，而且要有第一層'); reasons.append('遮第幾塊人工手動格式不對，沒有用')
    if mk: data['k'] = mk
    sv[c] = data
    vals = {'結構': s_auto, '第一層（電腦）': ' '.join(l1) if l1 else None, '第二層（電腦）': ' '.join(l2) if l2 else None,
            '第三層（開源最細）': ' '.join(l3) if l3 else None, '正式評量遮第幾塊（電腦）': mk_auto, '需要檢查': '；'.join(reasons) or None}
    for k, v in (('第一層', l1), ('第二層', l2), ('第三層', l3)):
        if v: stat[k] += 1
    stat['遮'] += bool(mk)
    src = ws.cell(r, 3)
    for h in H:
        cell = ws.cell(r, COL[h])
        if h in vals: cell.value = vals[h]
        cell.font, cell.alignment, cell.border = copy.copy(src.font), copy.copy(src.alignment), copy.copy(src.border)
        if h in TEACHER: cell.fill = green
        elif h in ('第一層（電腦）', '第二層（電腦）') and cell.value: cell.fill = yellow
        else: cell.fill = none
        if h in ('第一層狀態', '第二層狀態') and cell.value == '待討論': cell.fill = red
end = openpyxl.utils.get_column_letter(len(H))
ws.auto_filter.ref = f'A1:{end}{ws.max_row}'; ws.freeze_panes = 'B2'
ws.data_validations.dataValidation = []
for h, items in (('老師修正：結構', STRUCTS), ('第一層狀態', STATUS), ('第二層狀態', STATUS), ('正式評量遮第幾塊（人工手動）', ['1', '2', '3'])):
    dv = DataValidation(type='list', formula1='"' + ','.join(items) + '"', allow_blank=True)
    dv.add(f'{openpyxl.utils.get_column_letter(COL[h])}2:{openpyxl.utils.get_column_letter(COL[h])}{ws.max_row}'); ws.add_data_validation(dv)
# 讀不出聲音的部件：加進「部首變體對照」頁當候選（「對照成」留空，老師填了才生效；已經在表裡的只更新說明）
known = {c for r in range(2, vs.max_row + 1) for c in chars_of(vs.cell(r, 1).value or '')}
for p, n in sorted(unread.items(), key=lambda x: -x[1]):
    desc = f'電腦找到：第一層有 {n} 個字讀不出聲音（例：{"".join(unread_ex[p][:6])}；U+{ord(p[0]):04X}）。請在「對照成」填有讀音的字，沒有就留空'
    if p not in known: vs.append([p, None, desc, None])
    else:
        for r in range(2, vs.max_row + 1):
            if p in chars_of(vs.cell(r, 1).value or '') and not vs.cell(r, 2).value: vs.cell(r, 3).value = desc
wb.save(TABLE)
print(f'寫入 {len(sv)} 列；有第一層 {stat["第一層"]}、第二層 {stat["第二層"]}、第三層 {stat["第三層"]}、有遮擋規則 {stat["遮"]}')
if unread: print('第一層讀不出聲音的部件（要不要加進「部首變體對照」頁）：', sorted(unread.items(), key=lambda x: -x[1])[:25])

# 7. 產生 shared/parts-layers.js
head = '''/* 部件三層資料（自動產生，請不要手改；老師改部件表 docs/materials/parts-table.xlsx 的綠色欄位，再由 docs/scripts/make_parts_columns.py 重新產生）。
   設計見 docs/design/parts-layers.md、parts-layers-impl.md。步驟 3a 起平台只透過 window.Parts 讀這個檔，現在平台還沒讀。
   每個字：s＝結構（老師修正優先）；1、2、3＝第一～第三層，a＝電腦拆法、m＝老師的人工手動拆法（只有狀態是已確認或待確認才寫進來）、
   t＝人工手動的狀態（已確認／待確認／待討論）；k＝正式評量遮第幾塊（按第一層的順序數，沒有規則的字沒有 k）。
   部件用半形空格分開，同一部件出現幾次寫幾次；足、酋已當成一整塊。 */
'''
body = 'const PARTS_LAYERS = {\n' + ',\n'.join(f'{json.dumps(c, ensure_ascii=False)}:{json.dumps(d, ensure_ascii=False, separators=(",", ":"))}' for c, d in sv.items()) + '\n};\n'
open(OUT, 'w', encoding='utf8').write(head + body)
print(OUT, len((head + body).encode('utf8')) // 1024, 'KB')

# 8. 過渡：已確認的人工手動拆法寫進 lessons.js 的 PARTS_A、PARTS_B，平台目前的讀法不變
bridge = {'PARTS_A': {}, 'PARTS_B': {}}
for c, d in sv.items():
    for n, name in ((1, 'PARTS_A'), (2, 'PARTS_B')):
        x = d.get(str(n), {})
        if x.get('t') == '已確認' and x.get('m'): bridge[name][c] = x['m']
s = open(LES, encoding='utf8').read(); before = s
for name, d in bridge.items():
    line = f'const {name} = ' + json.dumps(d, ensure_ascii=False, separators=(',', ':')) + ';'
    s, n = re.subn(rf'^const {name} = .*;$', lambda m: line, s, flags=re.M)
    if n != 1: sys.exit(f'lessons.js 找不到 const {name} 那一行')
if s != before: open(LES, 'w', encoding='utf8').write(s); print('lessons.js 已更新：PARTS_A', len(bridge['PARTS_A']), '字、PARTS_B', len(bridge['PARTS_B']), '字；記得跑 node docs/scripts/bump.js 換版本號')
else: print('lessons.js 的 PARTS_A、PARTS_B 沒有變動：', len(bridge['PARTS_A']), '字、', len(bridge['PARTS_B']), '字')
for b in bad: print('格式不對、沒有寫進平台：', b)
